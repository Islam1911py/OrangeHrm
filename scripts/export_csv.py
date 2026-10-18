"""Regenerate docs/test-cases.csv from docs/test-cases.xlsx (needs: pip install openpyxl).
The CSV is the plain-text copy for AI tools and git diffs. Edit the xlsx, then run:  python scripts/export_csv.py
Execution columns (Status, Actual Result, Tester, Date, Defect) and the generated Automation column are not exported."""
import csv
import openpyxl

KEEP = ["TC ID", "Section", "Test Title", "Test Type", "Test Approach", "Priority", "Preconditions",
        "Test Data", "Steps", "Expected Result", "5.9 Check", "Notes"]
ws = openpyxl.load_workbook("docs/test-cases.xlsx")["Test Cases"]
header = [c.value for c in ws[1]]
idx = [header.index(h) for h in KEEP]
with open("docs/test-cases.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(KEEP)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            w.writerow(["" if row[i] is None else row[i] for i in idx])
print("wrote docs/test-cases.csv")
