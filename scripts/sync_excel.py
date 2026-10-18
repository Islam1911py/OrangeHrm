import json
import os
import openpyxl

json_file = 'test-results/results.json'
excel_file = 'docs/test-cases.xlsx'

if not os.path.exists(json_file):
    print("❌ ملف test-results/results.json مش موجود. شغّل التستات الأول بـ npx playwright test")
    exit(1)

if not os.path.exists(excel_file):
    print("❌ ملف docs/test-cases.xlsx مش موجود!")
    exit(1)

# 1. قراءة نتائج Playwright
with open(json_file, 'r', encoding='utf-8') as f:
    res = json.load(f)

test_status = {}

def extract_results(suite):
    for spec in suite.get('specs', []):
        title = spec.get('title', '')
        # البحث عن TC ID زي TC-3.1-04
        import re
        match = re.search(r'(TC-[\w.]+-\d+)', title)
        if match:
            tc_id = match.group(1)
            ok = spec.get('ok', False)
            test_status[tc_id] = 'Passed' if ok else 'Failed'
            
    for sub_suite in suite.get('suites', []):
        extract_results(sub_suite)

extract_results(res)

# 2. فتح ملف الإكسيل وتحديث عمود Status
wb = openpyxl.load_workbook(excel_file)
ws = wb['Test Cases']

# معرفة مكان العمدان
header = [cell.value for cell in ws[1]]
tc_id_col = header.index('TC ID') + 1 if 'TC ID' in header else 1
status_col = header.index('Status') + 1 if 'Status' in header else None

if not status_col:
    print("❌ عمود Status مش موجود في شيت Test Cases!")
    exit(1)

updated_count = 0
for row in range(2, ws.max_row + 1):
    tc_id = ws.cell(row=row, column=tc_id_col).value
    if tc_id:
        tc_id_str = str(tc_id).strip()
        if tc_id_str in test_status:
            ws.cell(row=row, column=status_col).value = test_status[tc_id_str]
            updated_count += 1
        else:
            # لو التست لسه متعملش بنسيبه Not Run
            if not ws.cell(row=row, column=status_col).value:
                ws.cell(row=row, column=status_col).value = 'Not Run'

# 3. حفظ الملف
wb.save(excel_file)
print(f"✅ تم تحديث {updated_count} حالة اختبار بنجاح جوه {excel_file}!")
print("📊 افتح شيت Summary في الإكسيل هتلاقي كل الأرقام والنسب اتحدثت أوتوماتيك!")